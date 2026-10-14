import os, glob

UP = "/workspace/.uploads"
OUT = "/workspace/.extract"
os.makedirs(OUT, exist_ok=True)

files = sorted(glob.glob(os.path.join(UP, "*")))
report = []
for f in files:
    base = os.path.basename(f)
    ext = os.path.splitext(base)[1].lower()
    out_path = os.path.join(OUT, base + ".txt")
    text = ""
    try:
        if ext == ".pdf":
            import pdfplumber
            parts = []
            with pdfplumber.open(f) as pdf:
                for i, page in enumerate(pdf.pages):
                    t = page.extract_text() or ""
                    parts.append(f"\n===== [page {i+1}] =====\n" + t)
            text = "\n".join(parts)
        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(f)
            parts = []
            for i, slide in enumerate(prs.slides):
                buf = [f"\n===== [slide {i+1}] ====="]
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        buf.append(shape.text_frame.text.strip())
                    if shape.has_table:
                        for row in shape.table.rows:
                            cells = [c.text.strip() for c in row.cells]
                            buf.append(" | ".join(cells))
                if slide.has_notes_slide:
                    nt = slide.notes_slide.notes_text_frame.text.strip()
                    if nt:
                        buf.append("[notes] " + nt)
                parts.append("\n".join(buf))
            text = "\n".join(parts)
    except Exception as e:
        text = f"[ERROR] {e}"
    with open(out_path, "w", encoding="utf-8") as fh:
        fh.write(text)
    report.append((base, len(text), out_path))

for base, n, p in report:
    print(f"{n:>8} chars | {base}")
